"""Congelar uma resposta (PEND-42): fica nem a conferir nem confirmada, com a categoria travada, até alguém descongelar.

Sem IA, sem nuvem e sem planilha: frame, respostas e classificação entram por dicionário.
Rodar:  python -m unittest tests.test_congelar -v
"""
from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="tabulador_congelar_")
os.environ["TABULADOR_OUTPUT"] = _TMP
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import codeframe as CF  # noqa: E402
import coding as CD  # noqa: E402
import eventos  # noqa: E402
import llm  # noqa: E402
import re  # noqa: E402
import supervisao as SUP  # noqa: E402


def _ia_falsa(system: str, user: str, schema: dict) -> dict:
    """IA falsa: manda tudo para a categoria 2 (nunca chama um provedor de verdade)."""
    return {"codificacoes": [{"rid": int(m.group(1)), "primaria": 2, "secundaria": None, "confianca": 0.9, "justificativa": "fake"}
                             for m in re.finditer(r'rid (\d+): "', user)]}


class TestCongelar(unittest.TestCase):
    def setUp(self):
        llm.set_cliente(_ia_falsa)
        self.addCleanup(llm.set_cliente, None)
        self.qid = f"Q_{self._testMethodName}"
        CF._gravar(self.qid, "frame.json", {
            "pergunta": self.qid, "status": "aprovado", "fixo": False, "versao": 1, "historico": [],
            "categorias": [{"codigo": 1, "nome": "Professores", "definicao": "x", "exemplos": [], "fixa": False},
                           {"codigo": 2, "nome": "Estrutura", "definicao": "y", "exemplos": [], "fixa": False},
                           *[dict(f) for f in CF.FIXAS]]})
        CF._gravar(self.qid, "respostas.json", {
            "pergunta": self.qid, "enunciado": "Teste", "n_respondeu": 4, "n_unicas": 4, "n_nsnr_auto": 0,
            "respostas": [{"rid": r, "texto": f"resposta {r}", "n": 1, "chave": f"resposta {r}", "nsnr": False, "respondent_ids": [r]} for r in (1, 2, 3, 4)]})
        CF._gravar(self.qid, "codificacao.json", {
            "pergunta": self.qid, "status": "rascunho",
            "itens": [{"rid": r, "primaria": 1, "secundaria": None, "confianca": 0.95, "origem": "llm",
                       "auditoria": {"por": "x", "ok": True, "em": "2026-10-02T10:00:00"}} for r in (1, 2, 3, 4)]}, permitir_reducao=True)

    def _item(self, rid):
        return next(i for i in CD.codificacao(self.qid)["itens"] if i["rid"] == rid)

    def test_congelada_nao_e_pendente_nem_confirmada(self):
        antes = CD.resumo_revisao(self.qid)
        CD.congelar(self.qid, 1, True, por="Sara")
        self.assertEqual(CD.situacao_revisao(self._item(1)), "congelada")
        r = CD.resumo_revisao(self.qid)
        self.assertEqual((r["n_congeladas"], r["n_confirmadas"], r["n_corrigidas"], r["n_revisadas"]), (1, 0, 0, 0))
        self.assertEqual(r["n_codificadas"], antes["n_codificadas"])  # continua classificada
        self.assertEqual(self._item(1)["congelada"]["por"], "Sara")
        self.assertEqual(CD.tabela(self.qid).set_index("rid").loc[1, "revisao"], "congelada")

    def test_congelada_recusa_qualquer_mudanca_ate_descongelar(self):
        CD.congelar(self.qid, 1)
        for kw in ({"validado": True}, {"primaria": 2}, {"secundaria": 2}, {"comentario": "x"}):
            with self.assertRaises(CD.RespostaCongelada):
                CD.atualizar_item(self.qid, 1, **kw)
        self.assertEqual(self._item(1)["primaria"], 1)
        CD.congelar(self.qid, 1, False)
        self.assertEqual(CD.situacao_revisao(self._item(1)), "pendente")  # volta a ser a conferir
        CD.atualizar_item(self.qid, 1, validado=True)
        self.assertEqual(CD.situacao_revisao(self._item(1)), "confirmada")

    def test_fora_das_operacoes_em_massa(self):
        CD.congelar(self.qid, 1)
        CD.congelar(self.qid, 2)
        self.assertEqual(CD.validar(self.qid, [1, 2, 3, 4]), 2)  # só as 3 e 4
        self.assertEqual(SUP.auto_aceitar(self.qid, 0.5)["aceitas"], 0)  # as 3 e 4 já estão confirmadas; 1 e 2 congeladas
        self.assertEqual({CD.situacao_revisao(self._item(r)) for r in (1, 2)}, {"congelada"})
        self.assertEqual(CD.resumo_revisao(self.qid)["n_confirmadas"], 2)

    def test_auto_aceitar_pula_congelada(self):
        CD.congelar(self.qid, 1)
        r = SUP.auto_aceitar(self.qid, 0.5)
        self.assertEqual(r["aceitas"], 3)
        self.assertEqual(r["ficaram"].get("congelada"), 1)

    def test_so_congela_o_que_ainda_nao_foi_decidido(self):
        CD.validar(self.qid, [2])
        CD.atualizar_item(self.qid, 3, primaria=2)
        for rid in (2, 3):
            with self.assertRaises(ValueError):
                CD.congelar(self.qid, rid)
        with self.assertRaises(ValueError):
            CD.congelar(self.qid, 99)  # não existe / não classificada

    def test_reclassificar_nao_mexe_na_congelada(self):
        CD.congelar(self.qid, 1)
        cod = CD.codificar(self.qid, forcar=True, preservar_humano=False, somente_faltantes=True)
        self.assertEqual(next(i for i in cod["itens"] if i["rid"] == 1)["congelada"]["por"], eventos.quem())
        ok = CD.recodificar(self.qid, rids=[1, 2])  # pedido explícito: a congelada fica de fora
        por_rid = {i["rid"]: i for i in ok["itens"]}
        self.assertTrue(por_rid[1].get("congelada"))
        self.assertEqual(por_rid[1]["primaria"], 1)  # a IA falsa mandaria para 2: não mexeu
        self.assertEqual(por_rid[2]["primaria"], 2)  # a não congelada foi refeita

    def test_planilha_de_revisao_nao_vaza_o_dicionario_e_ignora_a_congelada(self):
        import openpyxl
        CD.congelar(self.qid, 1, por="Sara")
        p = CD.exportar_revisao(self.qid)
        wb = openpyxl.load_workbook(p)
        ws = wb["revisao"]
        cab = {c.value: c.column for c in ws[1]}
        valores = {ws.cell(r, cab["rid"]).value: ws.cell(r, cab["congelada"]).value for r in range(2, ws.max_row + 1)}
        self.assertEqual(valores[1], "sim")  # legível, não o dicionário técnico
        self.assertIn(valores[2], ("", None))
        for r in range(2, ws.max_row + 1):  # a pessoa mexe nas linhas 1 (congelada) e 2
            if ws.cell(r, cab["rid"]).value in (1, 2):
                ws.cell(r, cab["REVISAO_PRIMARIA"]).value = "Estrutura"
        wb.save(p)
        cod = CD.importar_revisao(self.qid)
        por_rid = {i["rid"]: i for i in cod["itens"]}
        self.assertEqual(por_rid[1]["primaria"], 1)  # congelada: ignorada
        self.assertEqual(por_rid[2]["primaria"], 2)
        self.assertEqual((cod["_alteracoes"], cod["_congeladas_ignoradas"]), (1, 1))

    def test_log_de_eventos(self):
        CD.congelar(self.qid, 1, por="Sara")
        CD.congelar(self.qid, 1, False, por="João")
        ev = [(e["acao"], e["quem"], e["alvo"]) for e in eventos.listar(self.qid)]
        self.assertIn(("congelar", "Sara", 1), ev)
        self.assertIn(("descongelar", "João", 1), ev)


if __name__ == "__main__":
    unittest.main()
