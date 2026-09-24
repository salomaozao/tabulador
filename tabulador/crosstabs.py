"""Tabelas de cruzamento (banners) com % vertical.

Para cada pergunta, uma aba no Excel com: Total + níveis de cada banner nas colunas; nas linhas as
categorias (ou itens, para múltiplas; ou categorias do frame, para abertas codificadas; ou
distribuição/média, para numéricas). Percentuais sempre sobre a BASE da coluna (respondentes
válidos para a pergunta naquele segmento). Múltiplas e codificadas com secundária podem somar >100%.

Além das abas por pergunta, a aba "Geral" junta todas as variáveis numa única tabela (Variável |
Opção de resposta | Total | bandas dos banners) — ver `_escrever_aba_geral`. Os banners usados são
V.BANNERS_PADRAO por padrão, mas `gerar(banners=[...])` aceita uma lista escolhida na interface.

Saída: output/cruzamentos/cruzamentos.xlsx
"""
from __future__ import annotations

from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
import load
import variables as V

MIN_BASE = 10  # abaixo disso a base da coluna é destacada em vermelho


# ----------------------------------------------------------------------------- helpers
def _colunas_banner(df: pd.DataFrame, banners: list[str]) -> list[tuple[str, str, pd.Series]]:
    """[(banner, nível, máscara)] começando por Total."""
    cols = [("Total", "Total", pd.Series(True, index=df.index))]
    for b in banners:
        if b not in df.columns:
            raise KeyError(f"banner {b} não existe na base")
        niveis = V.DERIVADAS.get(b, {}).get("niveis") or sorted(df[b].dropna().unique().tolist())
        for n in niveis:
            cols.append((b, str(n), df[b].eq(n)))
    return cols


def _rotulo_banner(b: str) -> str:
    return V.DERIVADAS.get(b, {}).get("rotulo", b)


def _perguntas_tabulaveis(registro: dict) -> list[dict]:
    """Blocos a tabular: fechadas, múltiplas, numéricas, grades, abertas codificadas e derivadas."""
    blocos = []
    vistos = set()
    for var, ent in registro.items():
        tipo = ent.get("tipo")
        q = ent.get("pergunta") or var
        if tipo == "unica":
            blocos.append({"qid": q, "var": var, "tipo": "unica", "rotulo": ent["rotulo"], "enunciado": ent.get("enunciado"), "base": ent.get("base") or "todos", "niveis": ent.get("niveis")})
        elif tipo == "dicotomica" and q not in vistos:
            vistos.add(q)
            itens = [(v, e["opcao"]) for v, e in registro.items() if e.get("pergunta") == q and e.get("tipo") == "dicotomica"]
            blocos.append({"qid": q, "tipo": "multipla", "itens": itens, "rotulo": V.por_id(q)["rotulo"], "enunciado": ent.get("enunciado"), "base": ent.get("base") or "todos"})
        elif tipo == "numerica" and ent.get("opcao") is None:
            blocos.append({"qid": q, "var": var, "tipo": "numerica", "rotulo": ent["rotulo"], "enunciado": ent.get("enunciado"), "base": ent.get("base") or "todos"})
        elif tipo == "numerica" and q not in vistos:
            vistos.add(q)
            itens = [(v, e["opcao"]) for v, e in registro.items() if e.get("pergunta") == q and e.get("tipo") == "numerica"]
            blocos.append({"qid": q, "tipo": "grade", "itens": itens, "rotulo": V.por_id(q)["rotulo"], "enunciado": ent.get("enunciado"), "base": ent.get("base") or "todos"})
        elif tipo == "codificada" and var.endswith("_COD1"):
            blocos.append({"qid": q, "var": var, "tipo": "codificada", "rotulo": V.por_id(q)["rotulo"], "enunciado": ent.get("enunciado"), "base": ent.get("base") or "todos", "frame": ent.get("frame")})
        elif tipo == "derivada":
            blocos.append({"qid": var, "var": var, "tipo": "unica", "rotulo": ent["rotulo"], "enunciado": None, "base": "todos", "niveis": ent.get("niveis")})
    return blocos


def _eh_nps(qid: str) -> bool:
    try:
        return bool(V.por_id(qid).get("nps"))
    except KeyError:
        return False


def _nan(v) -> bool:
    return v is None or (isinstance(v, float) and np.isnan(v))


# ----------------------------------------------------------------------------- tabulação
def tabular(df: pd.DataFrame, bloco: dict, banners: list[str]) -> dict:
    """{'tipo', 'linhas', 'colunas': [(banner, nível)], 'base': [n], 'n': matriz, 'pct': matriz}"""
    cols = _colunas_banner(df, banners)
    base_masc = V.FILTROS[bloco["base"]](df)
    tipo = bloco["tipo"]
    linhas: list[str] = []
    n_mat: list[list] = []
    pct_mat: list[list] = []
    extra: dict = {}

    def _linha(rot, hit, valido, bases):
        linhas.append(rot)
        ns = [int((valido & m & hit).sum()) for _, _, m in cols]
        n_mat.append(ns)
        pct_mat.append([n / b if b else np.nan for n, b in zip(ns, bases)])

    if tipo == "unica":
        s = df[bloco["var"]]
        valido = base_masc & s.notna()
        freq = s[valido].value_counts().index.tolist()
        niveis = list(bloco.get("niveis") or [])
        if niveis:
            niveis += [v for v in freq if v not in niveis]  # valores inesperados vão para o fim
        else:
            # ordem de frequência; 'Outro. Qual?' e NS/NR vão para o fim
            niveis = sorted(freq, key=lambda v: (V.norm(str(v)).startswith("outr") or "ns/nr" in V.norm(str(v)), freq.index(v)))
        bases = [int((valido & m).sum()) for _, _, m in cols]
        for niv in niveis:
            _linha(str(niv), s.eq(niv), valido, bases)
    elif tipo == "multipla":
        itens = bloco["itens"]
        valido = base_masc & df[[v for v, _ in itens]].notna().any(axis=1)
        bases = [int((valido & m).sum()) for _, _, m in cols]
        for var, rot in itens:
            _linha(rot, df[var].eq(1), valido, bases)
        extra["nota"] = "Múltipla escolha: % de menções sobre a base; a soma pode passar de 100%."
    elif tipo == "codificada":
        c1 = bloco["var"]
        c2 = c1.replace("_COD1", "_COD2")
        valido = base_masc & df[c1].notna()
        bases = [int((valido & m).sum()) for _, _, m in cols]
        for cat in bloco["frame"]:
            k = cat["codigo"]
            _linha(f"{k}. {cat['nome']}", df[c1].eq(k) | df[c2].eq(k), valido, bases)
        extra["nota"] = "Aberta codificada: menções (primária + secundária) sobre a base; a soma pode passar de 100%."
    elif tipo == "numerica":
        s = pd.to_numeric(df[bloco["var"]], errors="coerce")
        valido = base_masc & s.notna()
        bases = [int((valido & m).sum()) for _, _, m in cols]
        for v in sorted(s[valido].unique().tolist()):
            _linha(f"{v:g}", s.eq(v), valido, bases)
        if _eh_nps(bloco["qid"]):
            for rot, hit in (("Promotores (9-10)", s.between(9, 10)), ("Neutros (7-8)", s.between(7, 8)), ("Detratores (0-6)", s.between(0, 6))):
                _linha(rot, hit, valido, bases)
            linhas.append("NPS (% promotores − % detratores) x100")
            n_mat.append([None] * len(cols))
            pct_mat.append([100 * (pct_mat[-3][i] - pct_mat[-1][i]) if bases[i] else np.nan for i in range(len(cols))])
        linhas.append("Média")
        n_mat.append([None] * len(cols))
        pct_mat.append([float(s[valido & m].mean()) if b else np.nan for (_, _, m), b in zip(cols, bases)])
        extra["linhas_valor"] = {"Média", "NPS (% promotores − % detratores) x100"}
    elif tipo == "grade":
        itens = bloco["itens"]
        valido = base_masc & df[[v for v, _ in itens]].notna().any(axis=1)
        bases = [int((valido & m).sum()) for _, _, m in cols]
        for var, rot in itens:
            s = pd.to_numeric(df[var], errors="coerce")
            linhas.append(f"{rot} — média")
            ns = [int((valido & m & s.notna()).sum()) for _, _, m in cols]
            n_mat.append(ns)
            pct_mat.append([float(s[valido & m].mean()) if n else np.nan for (_, _, m), n in zip(cols, ns)])
        extra["linhas_valor"] = set(linhas)
        extra["nota"] = "Grade 0-10: média por segmento; o bloco de n mostra quantos avaliaram cada atributo."
    else:
        raise ValueError(tipo)

    return {"tipo": tipo, "linhas": linhas, "colunas": [(b, n) for b, n, _ in cols], "base": bases, "n": n_mat, "pct": pct_mat, **extra}


# ----------------------------------------------------------------------------- excel
_FINO = Side(style="thin", color="BBBBBB")
_BORDA = Border(top=_FINO, bottom=_FINO, left=_FINO, right=_FINO)
_CAB = PatternFill("solid", fgColor="DDE6F2")
_NEG = Font(bold=True)
_CINZA = Font(italic=True, color="555555")


def _escrever_aba(ws, bloco: dict, tab: dict) -> None:
    ws["A1"] = f"{bloco['qid']} · {bloco['rotulo']}"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = bloco.get("enunciado") or ""
    ws["A3"] = f"Base: {V.FILTROS_DESCRICAO.get(bloco['base'], bloco['base'])}" + (f" · {tab['nota']}" if tab.get("nota") else "")
    ws["A3"].font = _CINZA
    r0, c = 5, 2
    colunas = tab["colunas"]
    i = 0
    while i < len(colunas):  # cabeçalho: banner (mesclado) / nível
        b = colunas[i][0]
        j = i
        while j < len(colunas) and colunas[j][0] == b:
            j += 1
        cel = ws.cell(row=r0, column=c + i, value=_rotulo_banner(b) if b != "Total" else "Total")
        cel.font = _NEG
        cel.alignment = Alignment(horizontal="center", wrap_text=True)
        if j - i > 1:
            ws.merge_cells(start_row=r0, start_column=c + i, end_row=r0, end_column=c + j - 1)
        for k in range(i, j):
            cel = ws.cell(row=r0 + 1, column=c + k, value=colunas[k][1])
            cel.font, cel.fill, cel.border = _NEG, _CAB, _BORDA
            cel.alignment = Alignment(horizontal="center", wrap_text=True, vertical="top")
        i = j
    ws.cell(row=r0 + 1, column=1).fill = _CAB
    linha = r0 + 2
    ws.cell(row=linha, column=1, value="Base (n)").font = _NEG
    for k, b in enumerate(tab["base"]):
        cel = ws.cell(row=linha, column=c + k, value=b)
        cel.border = _BORDA
        cel.font = Font(bold=True, color="B00020") if 0 < b < MIN_BASE else _NEG
    linha += 1
    linhas_valor = tab.get("linhas_valor", set())
    ws.cell(row=linha, column=1, value="Média" if tab["tipo"] == "grade" else "% (coluna)").font = _CINZA
    linha += 1
    for rot, pcts in zip(tab["linhas"], tab["pct"]):
        ws.cell(row=linha, column=1, value=rot).border = _BORDA
        for k, v in enumerate(pcts):
            cel = ws.cell(row=linha, column=c + k, value=None if _nan(v) else float(v))
            cel.border = _BORDA
            cel.number_format = "0.0" if rot in linhas_valor else "0%"
        linha += 1
    linha += 1
    ws.cell(row=linha, column=1, value="n (avaliaram)" if tab["tipo"] == "grade" else "n").font = _CINZA
    linha += 1
    for rot, ns in zip(tab["linhas"], tab["n"]):
        if all(v is None for v in ns):
            continue
        ws.cell(row=linha, column=1, value=rot).border = _BORDA
        for k, v in enumerate(ns):
            ws.cell(row=linha, column=c + k, value=v).border = _BORDA
        linha += 1
    ws.column_dimensions["A"].width = 48
    for k in range(len(colunas)):
        ws.column_dimensions[get_column_letter(c + k)].width = 13
    ws.freeze_panes = ws.cell(row=r0 + 2, column=2)


def _escrever_aba_geral(ws, blocos_tabs: list[tuple[dict, dict]], banners: list[str]) -> None:
    """Uma única tabela com todas as variáveis: Variável | Opção de resposta | Total | bandas dos banners."""
    ws["A1"] = "Cruzamento geral — todas as variáveis"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = f"Banners: {', '.join(_rotulo_banner(b) for b in banners)}"
    ws["A2"].font = _CINZA
    colunas = blocos_tabs[0][1]["colunas"] if blocos_tabs else []
    r0, c = 4, 3
    i = 0
    while i < len(colunas):  # cabeçalho: banner (mesclado) / nível, a partir da coluna C
        b = colunas[i][0]
        j = i
        while j < len(colunas) and colunas[j][0] == b:
            j += 1
        cel = ws.cell(row=r0, column=c + i, value=_rotulo_banner(b) if b != "Total" else "Total")
        cel.font = _NEG
        cel.alignment = Alignment(horizontal="center", wrap_text=True)
        if j - i > 1:
            ws.merge_cells(start_row=r0, start_column=c + i, end_row=r0, end_column=c + j - 1)
        for k in range(i, j):
            cel = ws.cell(row=r0 + 1, column=c + k, value=colunas[k][1])
            cel.font, cel.fill, cel.border = _NEG, _CAB, _BORDA
            cel.alignment = Alignment(horizontal="center", wrap_text=True, vertical="top")
        i = j
    for col, rotulo in ((1, "Variável"), (2, "Opção de resposta")):
        cel = ws.cell(row=r0 + 1, column=col, value=rotulo)
        cel.font, cel.fill = _NEG, _CAB
    linha = r0 + 2
    for bloco, tab in blocos_tabs:
        linhas_valor = tab.get("linhas_valor", set())
        for rot, pcts in zip(tab["linhas"], tab["pct"]):
            ws.cell(row=linha, column=1, value=bloco["rotulo"]).border = _BORDA
            ws.cell(row=linha, column=2, value=rot).border = _BORDA
            for k, v in enumerate(pcts):
                cel = ws.cell(row=linha, column=c + k, value=None if _nan(v) else float(v))
                cel.border = _BORDA
                cel.number_format = "0.0" if rot in linhas_valor else "0%"
            linha += 1
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 42
    for k in range(len(colunas)):
        ws.column_dimensions[get_column_letter(c + k)].width = 13
    ws.freeze_panes = ws.cell(row=r0 + 2, column=3)


def gerar(banners: list[str] | None = None, verbose: bool = True):
    banners = banners or V.BANNERS_PADRAO
    df, registro = load.carregar_base()
    blocos = _perguntas_tabulaveis(registro)
    config.garantir_pastas()
    p = config.CRUZAMENTOS_OUT / "cruzamentos.xlsx"
    longo = []
    blocos_tabs = []
    with pd.ExcelWriter(p, engine="openpyxl") as xw:
        idx = pd.DataFrame([{"aba": b["qid"][:31], "pergunta": b["qid"], "tipo": b["tipo"], "rotulo": b["rotulo"],
                             "base": V.FILTROS_DESCRICAO.get(b["base"], b["base"])} for b in blocos])
        idx.to_excel(xw, sheet_name="Índice", index=False)
        ws0 = xw.sheets["Índice"]
        ws0.column_dimensions["A"].width = 26
        ws0.column_dimensions["D"].width = 60
        ws0.column_dimensions["E"].width = 50
        ws0.cell(row=len(idx) + 3, column=1, value=f"Banners: {', '.join(_rotulo_banner(b) for b in banners)} · gerado em {datetime.now():%d/%m/%Y %H:%M} · N={len(df)}")
        usados = set()
        for b in blocos:
            tab = tabular(df, b, banners)
            blocos_tabs.append((b, tab))
            nome = b["qid"][:31]
            while nome in usados:
                nome = nome[:29] + "_2"
            usados.add(nome)
            _escrever_aba(xw.book.create_sheet(nome), b, tab)
            for rot, ns, pcts in zip(tab["linhas"], tab["n"], tab["pct"]):
                for (bn, niv), base, n, pc in zip(tab["colunas"], tab["base"], ns, pcts):
                    longo.append({"pergunta": b["qid"], "tipo": b["tipo"], "linha": rot, "banner": bn, "nivel": niv,
                                  "base": base, "n": n, "valor": None if _nan(pc) else float(pc)})
        _escrever_aba_geral(xw.book.create_sheet("Geral", 1), blocos_tabs, banners)
        pd.DataFrame(longo).to_excel(xw, sheet_name="dados_longos", index=False)
    if verbose:
        print(f"  {len(blocos)} tabelas em {p}")
    return p


if __name__ == "__main__":
    gerar()
