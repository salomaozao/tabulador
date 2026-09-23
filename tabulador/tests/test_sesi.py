"""Testes do projeto SESI (planilha 'plana') e do loop de validação, sem chamar a OpenAI.

Rodar:  python -m unittest discover -s tests -v
Escreve numa pasta temporária (TABULADOR_OUTPUT), nunca em SESI_cat/output.
"""
from __future__ import annotations

import os
import re
import sys
import tempfile
import unittest
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="tabulador_test_")
os.environ["TABULADOR_OUTPUT"] = _TMP
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402

import codebook  # noqa: E402
import codeframe as CF  # noqa: E402
import coding as CD  # noqa: E402
import config  # noqa: E402
import crosstabs  # noqa: E402
import exportar  # noqa: E402
import llm  # noqa: E402
import load  # noqa: E402
import projetos  # noqa: E402
import variables as V  # noqa: E402

projetos.ativar("sesi")


def _llm_fake(system: str, user: str, schema: dict) -> dict:
    props = schema["properties"]
    if "categorias" in props and "codigo_anterior" in props["categorias"]["items"]["properties"]:
        # refino: funde 2 em 1, mantém 3, cria 'Transporte'
        return {"raciocinio": "fake refino", "categorias": [
            {"codigo_anterior": 1, "incorpora": [2], "nome": "Formatura/mudança", "definicao": "fim do ciclo ou mudança", "exemplos": []},
            {"codigo_anterior": 3, "incorpora": [], "nome": "Ensino", "definicao": "qualidade do ensino", "exemplos": []},
            {"codigo_anterior": None, "incorpora": [], "nome": "Transporte", "definicao": "transporte escolar", "exemplos": []},
        ]}
    if "categorias" in props:
        return {"raciocinio": "fake", "categorias": [
            {"nome": "Formatura", "definicao": "último ano / vai se formar", "exemplos": ["Forma -se este ano"]},
            {"nome": "Mudança", "definicao": "mudança de cidade", "exemplos": ["mudança"]},
            {"nome": "Ensino", "definicao": "qualidade do ensino/professores", "exemplos": ["professores"]},
        ]}
    saida = []
    for m in re.finditer(r'rid (\d+): "(.*?)"', user):
        rid, texto = int(m.group(1)), m.group(2).lower()
        prim = 1 if ("form" in texto or "último" in texto or "ultimo" in texto) else 2 if "mud" in texto else 3 if ("prof" in texto or "ensin" in texto) else 97
        saida.append({"rid": rid, "primaria": prim, "secundaria": None, "confianca": 0.9 if prim != 97 else 0.4, "justificativa": "fake"})
    return {"codificacoes": saida}


@unittest.skipUnless(projetos.pasta("sesi").joinpath("data", "base_processamento.xlsx").exists(), "planilha SESI não disponível")
class TestSesi(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        projetos.ativar("sesi")
        llm.set_cliente(_llm_fake)
        cls.df, cls.registro = load.executar(verbose=False)

    @classmethod
    def tearDownClass(cls):
        llm.set_cliente(None)

    def setUp(self):
        projetos.ativar("sesi")

    def test_1_load_plano(self):
        df = self.df
        self.assertTrue(str(config.OUTPUT_DIR).startswith(_TMP))
        self.assertEqual(len(df), 3743)
        self.assertEqual(df["respondent_id"].nunique(), len(df))
        for pii in ("ip_address", "email_address", "first_name", "last_name"):
            self.assertNotIn(pii, df.columns)
        self.assertNotIn("@", (config.BASE_OUT / "base.json").read_text(encoding="utf-8"))
        self.assertEqual(V.perguntas_codificaveis(), ["Q2", "Q4", "Q5", "Q29", "Q30", "Q32", "Q33", "Q35"])
        # nenhuma resposta aberta perdida pelos filtros de base
        bruto = __import__("pandas").read_excel(config.FONTE_XLSX, sheet_name="base")
        for q in V.perguntas_codificaveis():
            self.assertEqual(int(df[q].notna().sum()), int(bruto[q].notna().sum()), q)
        self.assertEqual((df["NPS_GRUPO"] == "Promotores (9-10)").sum(), df["Q3"].between(9, 10).sum())
        self.assertIn("estrutura física", self.registro["Q6"]["enunciado"])

    def test_2_loop_validacao(self):
        qid = "Q33"
        dados = CF.preparar(qid)
        self.assertEqual(dados["n_respondeu"], 84)
        f = CF.induzir(qid, forcar=True)
        self.assertIn("Formatura", [c["nome"] for c in f["categorias"]])
        CF.aprovar(qid)
        # amostra
        cod = CD.codificar(qid, limite=20)
        res = CD.resumo_revisao(qid)
        nsnr = sum(1 for r in dados["respostas"] if r["nsnr"])
        self.assertEqual(res["n_codificadas"], 20 + nsnr)
        self.assertEqual(res["n_faltantes"], dados["n_unicas"] - 20 - nsnr)
        with self.assertRaises(RuntimeError):
            CD.aprovar(qid)  # ainda faltam respostas
        # confirmar 2, corrigir 1
        rids = [i["rid"] for i in cod["itens"]]
        self.assertEqual(CD.validar(qid, rids[:2]), 2)
        alvo = next(i for i in cod["itens"] if i["rid"] not in rids[:2] and i["primaria"] != 2)
        CD.atualizar_item(qid, alvo["rid"], primaria=2)
        res = CD.resumo_revisao(qid)
        self.assertEqual((res["n_confirmadas"], res["n_corrigidas"], res["n_avaliadas_ia"]), (2, 1, 3))
        self.assertAlmostEqual(res["acerto_ia"], 2 / 3, places=3)
        # resposta fora da amostra: classificar à mão cria o item
        faltante = next(r["rid"] for r in dados["respostas"] if r["rid"] not in {i["rid"] for i in CD.codificacao(qid)["itens"]})
        with self.assertRaises(ValueError):
            CD.atualizar_item(qid, faltante, comentario="x")
        CD.atualizar_item(qid, faltante, primaria=3)
        self.assertEqual(CD.resumo_revisao(qid)["n_avaliadas_ia"], 3)  # não entra na taxa de acerto
        # reclassificar tudo preserva confirmadas/corrigidas
        CD.recodificar(qid, nao_revisados=True)
        itens = {i["rid"]: i for i in CD.codificacao(qid)["itens"]}
        self.assertTrue(all(itens[r].get("validado") for r in rids[:2]))
        cod = CD.codificar(qid, somente_faltantes=True)
        self.assertEqual(CD.resumo_revisao(qid)["n_faltantes"], 0)
        itens = {i["rid"]: i for i in cod["itens"]}
        self.assertTrue(itens[rids[0]].get("validado"))
        self.assertEqual(itens[alvo["rid"]]["primaria"], 2)
        self.assertEqual(itens[alvo["rid"]]["origem"], "humano")
        # refino do frame pela IA: 2 fundida em 1, nova categoria 4, códigos preservados
        f = CF.refinar(qid, "junte formatura e mudança; crie transporte")
        codigos = [c["codigo"] for c in f["categorias"]]
        self.assertEqual(codigos, [1, 3, 4, CF.CODIGO_OUTROS, CF.CODIGO_NSNR])
        self.assertFalse(any(i["primaria"] == 2 for i in CD.codificacao(qid)["itens"]))
        self.assertEqual(f["status"], "aprovado")  # já há classificação: editar não pede nova aprovação
        CF.desfazer_refino(qid)
        self.assertEqual([c["codigo"] for c in CF.frame(qid)["categorias"]], [1, 2, 3, 97, 98])
        CF.aprovar(qid)
        CD.aprovar(qid)
        CD.aplicar_na_base(verbose=False)
        df, _ = load.carregar_base()
        self.assertEqual(int(df["Q33_COD1"].notna().sum()), 84)

    def test_3_saidas(self):
        if (CD.codificacao("Q33") or {}).get("status") != "aprovado":
            self.test_2_loop_validacao()
        p = exportar.gerar(verbose=False)
        wb = openpyxl.load_workbook(p, read_only=True)
        ws = wb["base"]
        cab = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        i = cab.index("Q33_CAT")
        self.assertIn("Q33_CAT2", cab)
        valores = [r[i] for r in ws.iter_rows(min_row=2, values_only=True) if r[i] is not None]
        self.assertEqual(len(valores), 84)
        self.assertIn("Categorias", wb.sheetnames)
        wb.close()
        self.assertTrue(codebook.gerar(verbose=False).exists())
        self.assertTrue(crosstabs.gerar(verbose=False).exists())
        df, _ = load.carregar_base()
        t = crosstabs.tabular(df, {"qid": "Q3", "var": "Q3", "tipo": "numerica", "base": "todos"}, ["VINCULO"])
        self.assertIn("NPS (% promotores − % detratores) x100", t["linhas"])

    def test_4_app_api(self):
        import app as A
        c = A.app.test_client()
        s = c.get("/api/status").get_json()
        self.assertEqual(s["projeto"]["slug"], "sesi")
        self.assertEqual(len(s["perguntas"]), 8)
        self.assertTrue(c.get("/").status_code == 200)
        r = c.post("/api/pergunta/Q2/frame/induzir", json={})
        self.assertEqual(r.status_code, 200, r.get_json())
        c.post("/api/pergunta/Q2/frame/aprovar", json={})
        r = c.post("/api/pergunta/Q2/codificar", json={"limite": 5}).get_json()
        cat = next(x for x in r["frame"]["categorias"] if not x.get("fixa"))
        self.assertGreater(r["revisao"]["n_faltantes"], 0)
        rid = r["itens"][0]["rid"] if r["itens"][0]["primaria"] is not None else next(i["rid"] for i in r["itens"] if i["primaria"] is not None)
        r = c.post(f"/api/pergunta/Q2/item/{rid}", json={"validado": True}).get_json()
        self.assertTrue(r["item"]["validado"])
        # editar categorias depois de classificar não trava a classificação (não pede nova aprovação)
        r = c.post("/api/pergunta/Q2/frame/categoria", json={"codigo": cat["codigo"], "nome": "Nome novo"}).get_json()
        self.assertEqual(r["frame"]["status"], "aprovado")
        r = c.delete(f"/api/pergunta/Q2/frame/categoria/{cat['codigo']}").get_json()
        self.assertEqual(r["frame"]["status"], "aprovado")
        r = c.post("/api/pergunta/Q2/codificar", json={"restantes": True}).get_json()
        self.assertEqual(r["revisao"]["n_faltantes"], 0)
        self.assertEqual(c.get("/api/progresso").status_code, 200)
        self.assertEqual(c.post("/api/gerar/codebook").status_code, 200)
        self.assertEqual(c.get("/arquivo/codebook").status_code, 200)
        self.assertEqual(c.post("/api/gerar/xyz").status_code, 404)


if __name__ == "__main__":
    unittest.main()
