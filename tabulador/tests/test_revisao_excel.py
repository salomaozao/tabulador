"""Testes de exportação e importação de revisao.xlsx."""
from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="tabulador_test_revisao_")
os.environ["TABULADOR_OUTPUT"] = _TMP
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

import pandas as pd
import openpyxl

import codeframe as CF
import coding as CD
import config


class TestRevisaoExcel(unittest.TestCase):
    def setUp(self):
        self.qid = "Q_TESTE"
        self.pasta = CF.pasta(self.qid)
        self.pasta.mkdir(parents=True, exist_ok=True)
        # Cria respostas.json
        respostas = {
            "qid": self.qid,
            "respostas": [
                {"rid": 1, "texto": "Atendimento excelente", "n": 5, "nsnr": False, "respondent_ids": [101, 102]},
                {"rid": 2, "texto": "Demora na fila", "n": 3, "nsnr": False, "respondent_ids": [103]},
            ]
        }
        (self.pasta / "respostas.json").write_text(json.dumps(respostas), encoding="utf-8")
        # Cria frame.json
        frame = {
            "qid": self.qid,
            "status": "aprovado",
            "categorias": [
                {"codigo": 1, "nome": "Elogio ao atendimento", "definicao": "Elogios"},
                {"codigo": 2, "nome": "Crítica ao tempo de espera", "definicao": "Reclamações"},
                {"codigo": 97, "nome": "Outros", "definicao": "Outros"},
            ]
        }
        (self.pasta / "frame.json").write_text(json.dumps(frame), encoding="utf-8")
        # Cria codificacao.json
        cod = {
            "qid": self.qid,
            "status": "rascunho",
            "itens": [
                {"rid": 1, "primaria": 1, "secundaria": None, "confianca": 0.95, "justificativa": "Elogiou", "origem": "llm", "auditoria": {"ok": True, "primaria": 1}, "validado_por": "auto"},
                {"rid": 2, "primaria": 97, "secundaria": None, "confianca": 0.60, "justificativa": "Genérico", "origem": "llm", "auditoria": {"ok": False, "primaria": 2}, "validado_por": None},
            ]
        }
        (self.pasta / "codificacao.json").write_text(json.dumps(cod), encoding="utf-8")

    def test_exportar_e_importar_revisao(self):
        p = CD.exportar_revisao(self.qid)
        self.assertTrue(p.exists())
        
        # Lê o Excel gerado
        df = pd.read_excel(p, sheet_name="revisao")
        colunas = list(df.columns)
        
        # Garante que colunas técnicas internas foram limpas
        self.assertNotIn("nsnr_regra", colunas)
        self.assertNotIn("auditoria", colunas)
        self.assertNotIn("validado_por", colunas)
        self.assertNotIn("respondent_ids", colunas)
        
        # Garante colunas de revisão
        self.assertIn("REVISAO_PRIMARIA", colunas)
        self.assertIn("REVISAO_SECUNDARIA", colunas)
        self.assertIn("COMENTARIO", colunas)
        self.assertIn("sugestao_auditor", colunas)
        
        # Simula o analista preenchendo a linha 2 no Excel
        wb = openpyxl.load_workbook(p)
        ws = wb["revisao"]
        # Acha coluna rid e REVISAO_PRIMARIA
        cab = {cell.value: idx + 1 for idx, cell in enumerate(ws[1])}
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=cab["rid"]).value == 2:
                ws.cell(row=row, column=cab["REVISAO_PRIMARIA"], value=2)
                ws.cell(row=row, column=cab["COMENTARIO"], value="Ajustado manualmente")
        wb.save(p)
        
        # Importa de volta
        resultado = CD.importar_revisao(self.qid)
        self.assertEqual(resultado["_alteracoes"], 1)
        
        # Verifica se atualizou no JSON
        cod_novo = CD.codificacao(self.qid)
        item2 = next(i for i in cod_novo["itens"] if i["rid"] == 2)
        self.assertEqual(item2["primaria"], 2)
        self.assertEqual(item2["origem"], "humano")
        self.assertEqual(item2["confianca"], 1.0)
        self.assertEqual(item2["comentario"], "Ajustado manualmente")
