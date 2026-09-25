"""Criação de projetos novos: projeto.json, rascunho sem IA, validador e o gabarito do SESI.

Rodar:  python -m unittest discover -s tests -v
Tudo em pastas temporárias; projetos.json real não é alterado.
"""
from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

_TMP = tempfile.mkdtemp(prefix="tabulador_novoprojeto_")
os.environ["TABULADOR_OUTPUT"] = _TMP
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

import openpyxl  # noqa: E402
from openpyxl.styles import PatternFill  # noqa: E402

import config  # noqa: E402
import load  # noqa: E402
import projetos  # noqa: E402
import variables as V  # noqa: E402
from novo_projeto import perfil as PF  # noqa: E402
from novo_projeto import projeto_json  # noqa: E402
from novo_projeto import rascunho as R  # noqa: E402
from novo_projeto import registrar as REG  # noqa: E402
from novo_projeto.validar import arquivo_projeto, carregar_modulo, validar  # noqa: E402

ROXO, AMARELO = PatternFill("solid", fgColor="FF7030A0"), PatternFill("solid", fgColor="FFFFFF00")
CAB = ["respondent_id", "email_address", "Q1", "Q2", "Q2_CAT", "Q3", "Q4", "Q4_CAT", "Q5", "Q5_CAT"]


def planilha_exemplo(arq: Path) -> Path:
    """40 respondentes. Q1 vínculo; Q2 aberta só p/ 'Outro'; Q3 nota 0-10 (com um NS/NR);
    Q4 aberta p/ nota <= 6; Q5 aberta p/ todos, mas Q5_CAT sem destaque (não foi pedida)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "base"
    ws.append(CAB)
    for c in ws[1]:
        c.fill = ROXO
    ws["E1"].fill = AMARELO  # Q2_CAT
    ws["H1"].fill = AMARELO  # Q4_CAT
    vinculos = ["Mãe", "Pai", "Outro. Qual?", "Mãe"]
    temas = ["professores atenciosos", "estrutura boa", "falta comunicação com a família", "mensalidade cara",
             "segurança na entrada", "merenda ruim", "gosto dos esportes", "biblioteca pequena"]
    for i in range(40):
        v = vinculos[i % 4]
        nota = "NS/NR" if i == 39 else i % 11
        ws.append([1000 + i, f"pessoa{i}@mail.com", v, f"avó número {i}" if v.startswith("Outro") else None, None,
                   nota, f"{temas[i % 8]} {i}" if isinstance(nota, int) and nota <= 6 else None, None,
                   f"{temas[(i + 3) % 8]} e {temas[i % 8]} ({i})", None])
    wb.save(arq)
    return arq


def completar(pasta: Path, pedidas: list[str]) -> dict:
    arq = pasta / "projeto.json"
    d = json.loads(arq.read_text(encoding="utf-8"))
    d["CONTEXTO_PROJETO"] = "Pesquisa de teste com pais de alunos."
    d["CONFERENCIA"] = {"responsavel": "Teste", "data": "2026-09-23", "origem_pedido": "planilha marcada", "abertas_pedidas": pedidas}
    REG.salvar_json(arq, d)
    return d


class TestNovoProjeto(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix="np_"))
        self.fonte = planilha_exemplo(Path(tempfile.mkdtemp()) / "pesquisa.xlsx")
        self.pasta = self.dir / "Teste_cat"
        self.perfil, self.proj = REG.criar_rascunho(self.pasta, self.fonte, "Teste")

    def spec(self, qid):
        return next(p for p in self.proj["PERGUNTAS"] if p["id"] == qid)

    def test_rascunho_deduz_do_dados(self):
        self.assertTrue((self.pasta / "data" / "pesquisa.xlsx").exists())
        self.assertEqual(self.proj["PII_MATCH"], ["email_address"])
        self.assertEqual(self.spec("Q1")["tipo"], "unica")
        self.assertEqual(set(self.spec("Q1")["niveis"]), {"Mãe", "Pai", "Outro. Qual?"})
        self.assertTrue(self.spec("Q3").get("nps"))
        self.assertEqual(self.proj["FILTROS"][self.spec("Q2")["base"]]["em"], ["Outro. Qual?"])
        self.assertEqual(self.proj["FILTROS"][self.spec("Q4")["base"]]["entre"], [0, 6])
        self.assertTrue(self.spec("Q2")["codificar"] and self.spec("Q4")["codificar"])
        self.assertFalse(self.spec("Q5")["codificar"])  # Q5_CAT sem destaque
        self.assertEqual(self.proj["SAIDA_PLANILHA"]["colunas"], {"Q2": "Q2_CAT", "Q4": "Q4_CAT"})

    def test_rascunho_recusado_ate_completar(self):
        res = validar(self.pasta)
        self.assertFalse(res["valido"])
        self.assertTrue(any("CONTEXTO_PROJETO" in e for e in res["erros"]))
        completar(self.pasta, ["Q2", "Q4"])
        res = validar(self.pasta)
        self.assertTrue(res["valido"], res["erros"])

    def test_conferencia_diferente_do_marcado(self):
        completar(self.pasta, ["Q2", "Q4", "Q5"])
        res = validar(self.pasta)
        self.assertTrue(any("Pedidas, mas NÃO marcadas" in e and "Q5" in e for e in res["erros"]))

    def _alterar(self, func):
        d = completar(self.pasta, ["Q2", "Q4"])
        func(d)
        REG.salvar_json(self.pasta / "projeto.json", d)
        return validar(self.pasta)

    def test_opcao_redigitada_e_erro(self):
        def f(d):
            next(p for p in d["PERGUNTAS"] if p["id"] == "Q1")["niveis"] = ["Mãe", "Pai", "Outro, qual?"]
        res = self._alterar(f)
        self.assertTrue(any("fora das opções" in e for e in res["erros"]))

    def test_filtro_que_apaga_resposta_e_erro(self):
        def f(d):
            regra = d["FILTROS"][next(p for p in d["PERGUNTAS"] if p["id"] == "Q4")["base"]]
            regra["entre"], regra["ou_respondeu"] = [0, 3], []
        res = self._alterar(f)
        self.assertTrue(any("Q4" in e and "apagadas" in e for e in res["erros"]))

    def test_dado_pessoal_declarado_e_erro(self):
        def f(d):
            d["PERGUNTAS"].append({"id": "EMAIL", "coluna": "email_address", "tipo": "meta", "rotulo": "e-mail"})
        res = self._alterar(f)
        self.assertTrue(any("EMAIL" in e and "dado pessoal" in e for e in res["erros"]))

    def test_projeto_json_roda_no_pipeline(self):
        completar(self.pasta, ["Q2", "Q4"])
        reg = projetos._registro()
        orig = projetos._registro
        projetos._registro = lambda: {**reg, "projetos": {**reg["projetos"], "teste_np": str(self.pasta)}}
        try:
            projetos._modulos.pop("teste_np", None)
            projetos.ativar("teste_np")
            self.assertEqual(config.NOME_PROJETO, "Teste")
            df, _ = load.executar(verbose=False)
            self.assertEqual(int(df["Q4"].notna().sum()), 27)  # notas 0-6 entre i=0..38
            self.assertIn("NPS_GRUPO", df.columns)
            self.assertEqual(V.perguntas_codificaveis(), ["Q2", "Q4"])
        finally:
            projetos._registro = orig
            projetos._modulos.pop("teste_np", None)
            projetos.ativar("sesi")

    def test_filtro_declarativo(self):
        import pandas as pd
        f = projeto_json.compilar_filtro({"pergunta": "N", "entre": [0, 6], "ou_respondeu": "A"})
        df = pd.DataFrame({"N": [3, 9, 9, None], "A": [None, None, "x", None]})
        self.assertEqual(f(df).tolist(), [True, False, True, False])


_SESI = projetos.pasta("sesi")


@unittest.skipUnless((_SESI / "data" / "base_processamento.xlsx").exists(), "planilha SESI não disponível")
class TestGabaritoSESI(unittest.TestCase):
    """O rascunho sem IA comparado ao projeto.py do SESI (feito e conferido à mão)."""

    @classmethod
    def setUpClass(cls):
        cls.gold = carregar_modulo(arquivo_projeto(_SESI) or (_SESI / "projeto.py"))
        cls.perfil = PF.perfilar(_SESI / "data" / "base_processamento.xlsx")
        cls.proj = R.gerar(cls.perfil, _SESI, "SESI")
        cls.por_coluna = {p.get("coluna", p["id"]): p for p in cls.proj["PERGUNTAS"]}

    def test_tipos_e_opcoes(self):
        for g in self.gold.PERGUNTAS:
            col = g.get("coluna", g["id"])
            r = self.por_coluna[col]
            self.assertEqual(r["tipo"], g["tipo"], col)
            if g.get("niveis"):
                self.assertEqual(set(r["niveis"]), set(g["niveis"]), col)
            self.assertEqual(bool(r.get("nps")), bool(g.get("nps")), col)

    def test_dados_pessoais(self):
        self.assertEqual(set(self.proj["PII_MATCH"]), set(self.gold.PII_MATCH))

    def test_abertas_pedidas_pela_cor(self):
        # as 7 com _CAT amarelo; a Q30 (sem destaque) fica de fora — é a pendência PEND-02
        marcadas = {p["id"] for p in self.proj["PERGUNTAS"] if p.get("codificar")}
        self.assertEqual(marcadas, {"Q2", "Q4", "Q5", "Q29", "Q32", "Q33", "Q35"})

    def test_filtros_das_abertas(self):
        import pandas as pd
        bruto = PF.ler(_SESI / "data" / "base_processamento.xlsx", "base")
        df = pd.DataFrame({g["id"]: (PF.numero if g["tipo"] == "numerica" else PF.texto)(bruto[g.get("coluna", g["id"])])
                           for g in self.gold.PERGUNTAS})
        filtros = {k: projeto_json.compilar_filtro(v) for k, v in self.proj["FILTROS"].items()}
        for g in self.gold.PERGUNTAS:
            if g["tipo"] != "aberta" or g.get("base", "todos") == "todos":
                continue
            ouro = self.gold.FILTROS[g["base"]](df)
            nosso = filtros[self.por_coluna[g["id"]]["base"]](df)
            self.assertTrue((nosso | ~df[g["id"]].notna()).all(), f"{g['id']}: rascunho perderia respostas")
            self.assertGreaterEqual((ouro == nosso).mean(), 0.98, g["id"])


if __name__ == "__main__":
    unittest.main()


class TestAssistenteApp(unittest.TestCase):
    """Fluxo do assistente pela API do app (upload -> validar -> criar), com projetos.json temporário."""

    def setUp(self):
        import shutil
        import app as APP
        self.app = APP.app.test_client()
        self.arq_orig = projetos.ARQUIVO
        tmp = Path(tempfile.mkdtemp()) / "projetos.json"
        shutil.copy(projetos.ARQUIVO, tmp)
        projetos.ARQUIVO = tmp
        self.ativo = config.PROJETO
        self.fonte = planilha_exemplo(Path(tempfile.mkdtemp()) / "pesquisa.xlsx")
        self.pasta = Path(tempfile.mkdtemp()) / "Assistente_cat"

    def tearDown(self):
        projetos.ARQUIVO = self.arq_orig
        projetos._modulos.pop("assistente", None)
        projetos.ativar(self.ativo)

    def test_fluxo_completo(self):
        with open(self.fonte, "rb") as f:
            r = self.app.post("/api/novo/iniciar", data={"nome": "Assistente", "pasta": str(self.pasta), "planilha": (f, "pesquisa.xlsx")},
                              content_type="multipart/form-data")
        self.assertEqual(r.status_code, 200, r.get_json())
        d = r.get_json()
        self.assertEqual(d["projeto"]["CONTEXTO_PROJETO"], "")  # PREENCHER vira campo vazio na tela
        proj = d["projeto"]
        # criar sem completar: recusado
        r = self.app.post("/api/novo/criar", json={"pasta": d["pasta"], "projeto": proj})
        self.assertEqual(r.status_code, 400)
        self.assertNotIn("assistente", projetos._registro()["projetos"])
        proj["CONTEXTO_PROJETO"] = "Pesquisa de teste."
        proj["CONFERENCIA"] = {"responsavel": "T", "data": "2026-09-23", "origem_pedido": "teste", "abertas_pedidas": ["Q2", "Q4"]}
        v = self.app.post("/api/novo/validar", json={"pasta": d["pasta"], "projeto": proj}).get_json()
        self.assertTrue(v["valido"], v["erros"])
        r = self.app.post("/api/novo/criar", json={"pasta": d["pasta"], "projeto": proj, "slug": "assistente"})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(config.PROJETO, "assistente")
        st = self.app.get("/api/status").get_json()
        self.assertEqual([p["qid"] for p in st["perguntas"]], ["Q2", "Q4"])
        self.assertTrue(st["base"]["carregada"])


class TestProvedorCLI(unittest.TestCase):
    """Provedor 'programa instalado no computador': llm.chamar_json delega ao llm_cli (sem rodar o programa)."""

    def test_delega_e_registra(self):
        import llm
        import llm_cli
        antes = (config.OPENAI_PROVEDOR, config.OPENAI_MODEL, llm_cli.chamar)
        chamadas = []
        llm_cli.chamar = lambda prov, s, u, schema, modelo: (chamadas.append((prov, modelo)) or ('```json\n{"ok": true}\n```', {"total_tokens": 3}))
        config.OPENAI_PROVEDOR, config.OPENAI_MODEL = "cli_claude", "haiku"
        try:
            llm.set_cliente(None)
            self.assertTrue(config.usa_cli())
            self.assertLessEqual(config.paralelo(), 2)
            r = llm.chamar_json("s", "u", {"type": "object"}, nome="teste")
            self.assertEqual(r, {"ok": True})
            self.assertEqual(chamadas, [("cli_claude", "haiku")])
        finally:
            config.OPENAI_PROVEDOR, config.OPENAI_MODEL, llm_cli.chamar = antes

    def test_atalho_cmd_vira_exe(self):
        import llm_cli
        d = Path(tempfile.mkdtemp())
        (d / "node_modules" / "x" / "bin").mkdir(parents=True)
        exe = d / "node_modules" / "x" / "bin" / "prog.exe"
        exe.write_text("")
        (d / "prog.cmd").write_text("@ECHO off\n" + r'"%dp0%\node_modules\x\bin\prog.exe"   %*' + "\n")
        antigo = os.environ["PATH"]
        os.environ["PATH"] = str(d) + os.pathsep + antigo
        try:
            self.assertEqual(Path(llm_cli.executavel("prog")).resolve(), exe.resolve())
        finally:
            os.environ["PATH"] = antigo
